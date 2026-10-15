import re
import json
import statistics
from pathlib import Path
from collections import Counter
from typing import Optional, Union

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def get_shape_text(shape):
    """Extract text from text boxes, placeholders, and tables."""
    parts = []

    if getattr(shape, "has_text_frame", False):
        text = shape.text.strip()
        if text:
            parts.append(text)

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [
                cell.text.strip()
                for cell in row.cells
                if cell.text.strip()
            ]
            if cells:
                parts.append(" | ".join(cells))

    return "\n".join(parts).strip()


def word_count(text):
    """Count words in text."""
    return len(re.findall(r"\b[\w'-]+\b", text))


def sentence_count(text):
    """Estimate the number of sentences."""
    sentences = re.split(r"[.!?]+", text.strip())
    return sum(
        1 for sentence in sentences
        if sentence.strip()
    )


def get_title(slide):
    """Find the slide title using its title placeholder first."""
    for shape in slide.shapes:
        try:
            if (
                shape.is_placeholder
                and shape.placeholder_format.idx == 0
                and shape.text.strip()
            ):
                return shape.text.strip().splitlines()[0]
        except (AttributeError, ValueError):
            pass

    # Fallback: use the first nonempty text shape.
    for shape in slide.shapes:
        text = get_shape_text(shape)
        if text:
            return text.splitlines()[0][:150]

    return ""


def get_font_sizes(slide):
    """Collect explicitly specified font sizes."""
    sizes = []

    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size is not None:
                        sizes.append(
                            round(run.font.size.pt, 1)
                        )

        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.font.size is not None:
                                sizes.append(
                                    round(run.font.size.pt, 1)
                                )

    return sizes


def get_shape_counts(slide):
    """Count visual elements and shapes."""
    images = 0
    charts = 0
    tables = 0
    text_boxes = 0
    other_shapes = 0

    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            images += 1
        elif getattr(shape, "has_chart", False):
            charts += 1
        elif getattr(shape, "has_table", False):
            tables += 1
        elif getattr(shape, "has_text_frame", False):
            text_boxes += 1
        else:
            other_shapes += 1

    return {
        "images": images,
        "charts": charts,
        "tables": tables,
        "text_boxes": text_boxes,
        "other_shapes": other_shapes
    }


# ============================================================
# MAIN ANALYSIS FUNCTION
# ============================================================

def analyze_presentation(ppt_path: Union[str, Path], output_dir: Optional[Union[str, Path]] = None) -> dict:
    """
    Analyzes a PowerPoint presentation (.pptx) file.
    Extracts structure metrics, text density, font sizes, shapes,
    and produces a comprehensive scoring and recommendation report.
    Saves 'pptanalysis.json' and 'pptanalysis.txt' in output_dir.
    """
    ppt_path = Path(ppt_path)
    if not ppt_path.exists():
        raise FileNotFoundError(f"PowerPoint file not found at: {ppt_path.resolve()}")

    if output_dir is None:
        output_dir = Path("output/ppt_analysed")
    else:
        output_dir = Path(output_dir)

    output_dir.mkdir(parents=True, exist_ok=True)

    print(f"\n[Astra-PPT] Loading PowerPoint: {ppt_path.name}")
    presentation = Presentation(str(ppt_path))

    if len(presentation.slides) == 0:
        raise ValueError("The PowerPoint contains no slides.")

    slides_report = []
    all_text = []
    all_titles = []

    # Analyze each slide
    for number, slide in enumerate(presentation.slides, start=1):
        slide_blocks = []

        title = get_title(slide)
        font_sizes = get_font_sizes(slide)
        shape_counts = get_shape_counts(slide)

        for shape in slide.shapes:
            text = get_shape_text(shape)
            if text:
                slide_blocks.append(text)
                all_text.append(text)

        text = "\n".join(slide_blocks).strip()

        if title:
            all_titles.append(title)

        words = word_count(text)
        characters = len(text)
        sentences = sentence_count(text)

        average_words_per_sentence = (
            round(words / sentences, 2)
            if sentences else 0
        )

        findings = []
        suggestions = []

        # Base score: 100
        score = 100

        # 1. Title Analysis
        if not title:
            findings.append("No detectable slide title.")
            suggestions.append("Add a clear, descriptive slide title.")
            score -= 12

        # 2. Text Density
        if words > 100:
            findings.append(f"Text-heavy slide: {words} words.")
            suggestions.append("Reduce text and highlight the key message.")
            score -= 18
        elif words > 70:
            findings.append(f"Relatively dense slide: {words} words.")
            suggestions.append("Shorten paragraphs and use concise bullet points.")
            score -= 10

        # 3. Low Text Content
        if words < 5:
            findings.append("Very little extractable text.")
            suggestions.append("Check whether the slide communicates its purpose. A visual-only slide may be intentional.")
            score -= 5

        # 4. Font Readability
        if font_sizes:
            minimum_font = min(font_sizes)
            if minimum_font < 12:
                findings.append(f"Small explicitly sized text: {minimum_font} pt.")
                suggestions.append("Increase small fonts where possible.")
                score -= 12
            elif minimum_font < 14:
                findings.append(f"Some text is below 14 pt. Minimum: {minimum_font} pt.")
                suggestions.append("Check readability on a projector.")
                score -= 6
        else:
            findings.append("No explicit font sizes could be extracted.")

        # 5. Long Sentences
        if average_words_per_sentence > 25:
            findings.append("Long sentences may reduce readability.")
            suggestions.append("Break long sentences into shorter statements.")
            score -= 5

        # 6. Visual Elements
        if shape_counts["images"]:
            suggestions.append("Check image quality, relevance, alignment, and attribution.")
        if shape_counts["charts"]:
            suggestions.append("Check chart titles, labels, units, and readability.")
        if shape_counts["tables"]:
            suggestions.append("Check table readability and simplify dense cells.")

        # 7. Empty Slide Check
        if not text and not shape_counts["images"] \
                and not shape_counts["charts"] \
                and not shape_counts["tables"]:
            findings.append("Potentially empty slide.")
            suggestions.append("Add meaningful content or remove the slide.")
            score -= 10

        # 8. Duplicate Title Check
        previous_titles = [
            item.lower().strip()
            for item in all_titles[:-1]
        ]

        if title and title.lower().strip() in previous_titles:
            findings.append("Duplicate title found on another slide.")
            suggestions.append("Check whether the repeated title is intentional.")

        # Final Slide Score bounded 0-100
        score = max(0, min(100, score))

        if not findings:
            findings.append("No issues detected by the basic automated checks.")

        slides_report.append({
            "slide_number": number,
            "title": title or "(No title detected)",
            "word_count": words,
            "character_count": characters,
            "sentence_count": sentences,
            "average_words_per_sentence": average_words_per_sentence,
            "image_count": shape_counts["images"],
            "chart_count": shape_counts["charts"],
            "table_count": shape_counts["tables"],
            "text_box_count": shape_counts["text_boxes"],
            "other_shape_count": shape_counts["other_shapes"],
            "font_sizes_pt": font_sizes,
            "minimum_font_size_pt": min(font_sizes) if font_sizes else None,
            "maximum_font_size_pt": max(font_sizes) if font_sizes else None,
            "basic_structure_score": score,
            "findings": findings,
            "suggestions": suggestions,
            "extracted_text": text
        })

    # Repeated 4-word phrases
    phrase_counts = Counter()
    for text in all_text:
        tokens = re.findall(r"\b[a-zA-Z0-9'-]+\b", text.lower())
        for i in range(len(tokens) - 3):
            phrase = " ".join(tokens[i:i + 4])
            phrase_counts[phrase] += 1

    repeated_phrases = [
        {"phrase": phrase, "count": count}
        for phrase, count in phrase_counts.most_common(30)
        if count > 1
    ]

    # Presentation-level metrics
    slide_count = len(slides_report)
    total_words = sum(slide["word_count"] for slide in slides_report)
    total_characters = sum(slide["character_count"] for slide in slides_report)
    total_images = sum(slide["image_count"] for slide in slides_report)
    total_charts = sum(slide["chart_count"] for slide in slides_report)
    total_tables = sum(slide["table_count"] for slide in slides_report)

    average_words_per_slide = round(total_words / slide_count, 2) if slide_count else 0
    average_score = round(
        statistics.mean(slide["basic_structure_score"] for slide in slides_report),
        1
    ) if slides_report else 0

    untitled_slides = sum(slide["title"] == "(No title detected)" for slide in slides_report)
    dense_slides = sum(slide["word_count"] > 100 for slide in slides_report)
    low_text_slides = sum(slide["word_count"] < 5 for slide in slides_report)
    slides_below_60 = sum(slide["basic_structure_score"] < 60 for slide in slides_report)

    overall_findings = []
    if untitled_slides:
        overall_findings.append(f"{untitled_slides} slide(s) have no detectable title.")
    if dense_slides:
        overall_findings.append(f"{dense_slides} slide(s) contain more than 100 words.")
    if low_text_slides:
        overall_findings.append(f"{low_text_slides} slide(s) have fewer than 5 words.")
    if slides_below_60:
        overall_findings.append(f"{slides_below_60} slide(s) scored below 60/100.")
    if not overall_findings:
        overall_findings.append("No major issues detected by the basic automated checks.")

    # Create Report Object
    report = {
        "file_name": ppt_path.name,
        "slide_count": slide_count,
        "total_word_count": total_words,
        "total_character_count": total_characters,
        "average_words_per_slide": average_words_per_slide,
        "total_images": total_images,
        "total_charts": total_charts,
        "total_tables": total_tables,
        "average_basic_structure_score": average_score,
        "untitled_slides": untitled_slides,
        "dense_slides": dense_slides,
        "low_text_slides": low_text_slides,
        "slides_below_60": slides_below_60,
        "overall_findings": overall_findings,
        "score_explanation": (
            "Rule-based heuristic based on detectable titles, "
            "text density, explicitly sized small fonts, "
            "sentence length, and potentially empty slides. "
            "This is not an AI assessment of visual design, "
            "factual accuracy, or persuasiveness."
        ),
        "slides": slides_report,
        "repeated_four_word_phrases": repeated_phrases
    }

    # Save JSON Report in Backend
    json_path = output_dir / "pptanalysis.json"
    with open(json_path, "w", encoding="utf-8") as file:
        json.dump(report, file, indent=2, ensure_ascii=False)

    # Generate Readable Text Report
    lines = []
    def add_line(text=""):
        lines.append(str(text))

    add_line("=" * 70)
    add_line("                 POWERPOINT ANALYSIS REPORT")
    add_line("=" * 70)
    add_line(f"File name: {ppt_path.name}")
    add_line(f"Total slides: {slide_count}")
    add_line(f"Total extracted words: {total_words}")
    add_line(f"Average words per slide: {average_words_per_slide}")
    add_line(f"Average basic structure score: {average_score}/100")
    add_line()
    add_line("SCORING EXPLANATION")
    add_line("-" * 70)
    add_line(report["score_explanation"])
    add_line()
    add_line("OVERALL FINDINGS")
    add_line("-" * 70)
    for finding in overall_findings:
        add_line(f"- {finding}")
    add_line()
    add_line("SLIDE-BY-SLIDE ANALYSIS")
    add_line("=" * 70)
    for slide in slides_report:
        add_line()
        add_line(f"SLIDE {slide['slide_number']}: {slide['title']}")
        add_line("-" * 70)
        add_line(f"Basic structure score: {slide['basic_structure_score']}/100")
        add_line()
        add_line("FINDINGS:")
        for finding in slide["findings"]:
            add_line(f"- {finding}")
        add_line()
        add_line("RECOMMENDATIONS:")
        if slide["suggestions"]:
            for suggestion in slide["suggestions"]:
                add_line(f"- {suggestion}")
        else:
            add_line("No specific recommendations.")
    add_line()
    add_line("REPEATED FOUR-WORD PHRASES")
    add_line("=" * 70)
    if repeated_phrases:
        for item in repeated_phrases:
            add_line(f"- '{item['phrase']}': {item['count']} occurrences")
    else:
        add_line("No repeated four-word phrases detected.")
    add_line()
    add_line("=" * 70)
    add_line("END OF REPORT")
    add_line("=" * 70)

    # Save Text Report in Backend
    txt_path = output_dir / "pptanalysis.txt"
    with open(txt_path, "w", encoding="utf-8") as file:
        file.write("\n".join(lines))

    print(f"[Astra-PPT] Analysis completed successfully!")
    print(f"[Astra-PPT] JSON Report saved to: {json_path.resolve()}")
    print(f"[Astra-PPT] Text Report saved to: {txt_path.resolve()}")

    return report


# Alias for backward compatibility
analyze_pptx = analyze_presentation


if __name__ == "__main__":
    project_root = Path.cwd()
    raw_dir = project_root / "data" / "raw"
    output_dir = project_root / "output" / "ppt_analysed"

    print("\n" + "=" * 60)
    print("          POWERPOINT ANALYZER")
    print("=" * 60)

    ppt_files = list(raw_dir.glob("*.pptx"))
    if not ppt_files:
        raise FileNotFoundError(
            f"No .pptx files found in {raw_dir.resolve()}.\n"
            "Please drop a PowerPoint presentation into the data/raw/ folder."
        )

    target_ppt = ppt_files[0]
    print(f"\nAuto-detected PowerPoint: {target_ppt.name}")
    analyze_presentation(target_ppt, output_dir)