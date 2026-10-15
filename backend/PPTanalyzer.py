import json
import re
import statistics
from collections import Counter
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE


def get_shape_text(shape):
    parts = []
    if getattr(shape, "has_text_frame", False):
        text = shape.text.strip()
        if text:
            parts.append(text)
    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                parts.append(" | ".join(cells))
    return "\n".join(parts).strip()


def word_count(text):
    return len(re.findall(r"\b[\w'-]+\b", text))


def sentence_count(text):
    return sum(1 for sentence in re.split(r"[.!?]+", text.strip()) if sentence.strip())


def get_title(slide):
    for shape in slide.shapes:
        try:
            if shape.is_placeholder and shape.placeholder_format.idx == 0 and shape.text.strip():
                return shape.text.strip().splitlines()[0]
        except (AttributeError, ValueError):
            pass
    for shape in slide.shapes:
        text = get_shape_text(shape)
        if text:
            return text.splitlines()[0][:150]
    return ""


def get_font_sizes(slide):
    sizes = []
    for shape in slide.shapes:
        if getattr(shape, "has_text_frame", False):
            for paragraph in shape.text_frame.paragraphs:
                for run in paragraph.runs:
                    if run.font.size is not None:
                        sizes.append(round(run.font.size.pt, 1))
        if getattr(shape, "has_table", False):
            for row in shape.table.rows:
                for cell in row.cells:
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            if run.font.size is not None:
                                sizes.append(round(run.font.size.pt, 1))
    return sizes


def get_shape_counts(slide):
    counts = {"images": 0, "charts": 0, "tables": 0, "text_boxes": 0, "other_shapes": 0}
    for shape in slide.shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            counts["images"] += 1
        elif getattr(shape, "has_chart", False):
            counts["charts"] += 1
        elif getattr(shape, "has_table", False):
            counts["tables"] += 1
        elif getattr(shape, "has_text_frame", False):
            counts["text_boxes"] += 1
        else:
            counts["other_shapes"] += 1
    return counts


def analyze_pptx(pptx_path, output_dir=None):
    """Analyze one PPTX and return a JSON-serializable report."""
    pptx_path = Path(pptx_path)
    presentation = Presentation(str(pptx_path))
    if len(presentation.slides) == 0:
        raise ValueError("The PowerPoint contains no slides.")

    slides_report, all_text, all_titles = [], [], []
    for number, slide in enumerate(presentation.slides, start=1):
        slide_blocks = []
        title = get_title(slide)
        font_sizes = get_font_sizes(slide)
        shape_counts = get_shape_counts(slide)
        for shape in slide.shapes:
            text = get_shape_text(shape)
            if text:
                slide_blocks.append(text)
                all_text.append(text)
        text = "\n".join(slide_blocks).strip()
        if title:
            all_titles.append(title)
        words = word_count(text)
        characters = len(text)
        sentences = sentence_count(text)
        average_words_per_sentence = round(words / sentences, 2) if sentences else 0
        findings, suggestions, score = [], [], 100

        if not title:
            findings.append("No detectable slide title.")
            suggestions.append("Add a clear, descriptive slide title.")
            score -= 12
        if words > 100:
            findings.append(f"Text-heavy slide: {words} words.")
            suggestions.append("Reduce text and highlight the key message.")
            score -= 18
        elif words > 70:
            findings.append(f"Relatively dense slide: {words} words.")
            suggestions.append("Shorten paragraphs and use concise bullet points.")
            score -= 10
        if words < 5:
            findings.append("Very little extractable text.")
            suggestions.append("Check whether the slide communicates its purpose. A visual-only slide may be intentional.")
            score -= 5
        if font_sizes:
            minimum_font = min(font_sizes)
            if minimum_font < 12:
                findings.append(f"Small explicitly sized text: {minimum_font} pt.")
                suggestions.append("Increase small fonts where possible.")
                score -= 12
            elif minimum_font < 14:
                findings.append(f"Some text is below 14 pt. Minimum: {minimum_font} pt.")
                suggestions.append("Check readability on a projector.")
                score -= 6
        else:
            findings.append("No explicit font sizes could be extracted.")
        if average_words_per_sentence > 25:
            findings.append("Long sentences may reduce readability.")
            suggestions.append("Break long sentences into shorter statements.")
            score -= 5
        if shape_counts["images"]:
            suggestions.append("Check image quality, relevance, alignment, and attribution.")
        if shape_counts["charts"]:
            suggestions.append("Check chart titles, labels, units, and readability.")
        if shape_counts["tables"]:
            suggestions.append("Check table readability and simplify dense cells.")
        if not text and not shape_counts["images"] and not shape_counts["charts"] and not shape_counts["tables"]:
            findings.append("Potentially empty slide.")
            suggestions.append("Add meaningful content or remove the slide.")
            score -= 10
        previous_titles = [item.lower().strip() for item in all_titles[:-1]]
        if title and title.lower().strip() in previous_titles:
            findings.append("Duplicate title found on another slide.")
            suggestions.append("Check whether the repeated title is intentional.")
        score = max(0, min(100, score))
        if not findings:
            findings.append("No issues detected by the basic automated checks.")
        slides_report.append({
            "slide_number": number,
            "title": title or "(No title detected)",
            "word_count": words,
            "character_count": characters,
            "sentence_count": sentences,
            "average_words_per_sentence": average_words_per_sentence,
            "image_count": shape_counts["images"],
            "chart_count": shape_counts["charts"],
            "table_count": shape_counts["tables"],
            "text_box_count": shape_counts["text_boxes"],
            "other_shape_count": shape_counts["other_shapes"],
            "font_sizes_pt": font_sizes,
            "minimum_font_size_pt": min(font_sizes) if font_sizes else None,
            "maximum_font_size_pt": max(font_sizes) if font_sizes else None,
            "basic_structure_score": score,
            "findings": findings,
            "suggestions": suggestions,
            "extracted_text": text,
        })

    phrase_counts = Counter()
    for text in all_text:
        tokens = re.findall(r"\b[a-zA-Z0-9'-]+\b", text.lower())
        for i in range(len(tokens) - 3):
            phrase_counts[" ".join(tokens[i:i + 4])] += 1
    repeated_phrases = [{"phrase": phrase, "count": count} for phrase, count in phrase_counts.most_common(30) if count > 1]

    slide_count = len(slides_report)
    total_words = sum(slide["word_count"] for slide in slides_report)
    total_characters = sum(slide["character_count"] for slide in slides_report)
    total_images = sum(slide["image_count"] for slide in slides_report)
    total_charts = sum(slide["chart_count"] for slide in slides_report)
    total_tables = sum(slide["table_count"] for slide in slides_report)
    average_words_per_slide = round(total_words / slide_count, 2)
    average_score = round(statistics.mean(slide["basic_structure_score"] for slide in slides_report), 1)
    untitled_slides = sum(slide["title"] == "(No title detected)" for slide in slides_report)
    dense_slides = sum(slide["word_count"] > 100 for slide in slides_report)
    low_text_slides = sum(slide["word_count"] < 5 for slide in slides_report)
    slides_below_60 = sum(slide["basic_structure_score"] < 60 for slide in slides_report)
    overall_findings = []
    if untitled_slides:
        overall_findings.append(f"{untitled_slides} slide(s) have no detectable title.")
    if dense_slides:
        overall_findings.append(f"{dense_slides} slide(s) contain more than 100 words.")
    if low_text_slides:
        overall_findings.append(f"{low_text_slides} slide(s) have fewer than 5 words.")
    if slides_below_60:
        overall_findings.append(f"{slides_below_60} slide(s) scored below 60/100.")
    if not overall_findings:
        overall_findings.append("No major issues detected by the basic automated checks.")

    report = {
        "file_name": pptx_path.name,
        "slide_count": slide_count,
        "total_word_count": total_words,
        "total_character_count": total_characters,
        "average_words_per_slide": average_words_per_slide,
        "total_images": total_images,
        "total_charts": total_charts,
        "total_tables": total_tables,
        "average_basic_structure_score": average_score,
        "untitled_slides": untitled_slides,
        "dense_slides": dense_slides,
        "low_text_slides": low_text_slides,
        "slides_below_60": slides_below_60,
        "overall_findings": overall_findings,
        "score_explanation": "Rule-based heuristic based on detectable titles, text density, explicitly sized small fonts, sentence length, and potentially empty slides. This is not an AI assessment of visual design, factual accuracy, or persuasiveness.",
        "slides": slides_report,
        "repeated_four_word_phrases": repeated_phrases,
    }
    if output_dir:
        output_dir = Path(output_dir)
        output_dir.mkdir(parents=True, exist_ok=True)
        (output_dir / "pptanalysis.json").write_text(json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8")
        (output_dir / "pptanalysis.txt").write_text(make_text_report(report), encoding="utf-8")
    return report


def make_text_report(report):
    lines = ["=" * 70, "                 POWERPOINT ANALYSIS REPORT", "=" * 70,
             f"File name: {report['file_name']}", f"Total slides: {report['slide_count']}",
             f"Total extracted words: {report['total_word_count']}",
             f"Average words per slide: {report['average_words_per_slide']}",
             f"Average basic structure score: {report['average_basic_structure_score']}/100", "",
             "SCORING EXPLANATION", "-" * 70, report["score_explanation"], "",
             "OVERALL FINDINGS", "-" * 70]
    lines.extend(f"- {x}" for x in report["overall_findings"])
    lines.extend(["", "SLIDE-BY-SLIDE ANALYSIS", "=" * 70])
    for slide in report["slides"]:
        lines.extend(["", f"SLIDE {slide['slide_number']}: {slide['title']}", "-" * 70,
                      f"Basic structure score: {slide['basic_structure_score']}/100", "", "FINDINGS:"])
        lines.extend(f"- {x}" for x in slide["findings"])
        lines.extend(["", "RECOMMENDATIONS:"])
        lines.extend(f"- {x}" for x in slide["suggestions"] or ["No specific recommendations."])
    lines.extend(["", "REPEATED FOUR-WORD PHRASES", "=" * 70])
    lines.extend(f"- '{x['phrase']}': {x['count']} occurrences" for x in report["repeated_four_word_phrases"])
    if not report["repeated_four_word_phrases"]:
        lines.append("No repeated four-word phrases detected.")
    lines.extend(["", "=" * 70, "END OF REPORT", "=" * 70])
    return "\n".join(lines)
